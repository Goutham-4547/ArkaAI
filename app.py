import csv
import json
import os
import random
import re
import hashlib
import math
import time
import shutil
import subprocess
import tempfile
import threading
import textwrap
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

# The model backend is selected at runtime by model_router.py.
# These names remain as compatibility aliases for a few legacy UI paths.
MODEL_NAME = ""
FALLBACK_MODEL = ""

# ---------- Free local emotional-signal model ----------
# SamLowe/roberta-base-go_emotions-onnx (MIT licensed), INT8 ONNX.
# It is used only as a lightweight emotion signal; it is NOT a diagnosis.
EMOTION_MODEL_REPO = "SamLowe/roberta-base-go_emotions-onnx"
EMOTION_MODEL_FILE = "onnx/model_quantized.onnx"
EMOTION_LABELS = [
    "admiration", "amusement", "anger", "annoyance", "approval",
    "caring", "confusion", "curiosity", "desire", "disappointment",
    "disapproval", "disgust", "embarrassment", "excitement", "fear",
    "gratitude", "grief", "joy", "love", "nervousness", "optimism",
    "pride", "realization", "relief", "remorse", "sadness", "surprise",
    "neutral",
]
_emotion_session = None
_emotion_tokenizer = None
_emotion_lock = threading.Lock()
_emotion_failed = False


def _emotion_model_load():
    """Lazy-load the free INT8 GoEmotions model on first emotional check."""
    global _emotion_session, _emotion_tokenizer, _emotion_failed
    with _emotion_lock:
        if _emotion_session is not None and _emotion_tokenizer is not None:
            return _emotion_session, _emotion_tokenizer
        if _emotion_failed:
            return None, None
        try:
            import numpy as np  # noqa: F401
            import onnxruntime as ort
            from huggingface_hub import hf_hub_download
            from transformers import AutoTokenizer

            model_path = hf_hub_download(
                repo_id=EMOTION_MODEL_REPO,
                filename=EMOTION_MODEL_FILE,
            )
            tokenizer = AutoTokenizer.from_pretrained(
                EMOTION_MODEL_REPO,
                use_fast=True,
            )
            session = create_onnx_session(model_path)
            _emotion_session = session
            _emotion_tokenizer = tokenizer
            return session, tokenizer
        except Exception as exc:
            print("Emotion AI unavailable; using fallback emotion heuristics:", exc)
            _emotion_failed = True
            return None, None


def _emotion_fallback(text):
    """Small zero-dependency fallback so emotional support still works without downloads."""
    t = (text or "").lower()
    groups = {
        "sadness": ("sad", "depressed", "upset", "cry", "lonely", "hurt", "down"),
        "fear": ("scared", "afraid", "worried", "panic", "anxious", "nervous"),
        "anger": ("angry", "mad", "furious", "hate", "annoyed", "irritated"),
        "confusion": ("confused", "don't understand", "dont understand", "lost", "what does this mean"),
        "annoyance": ("frustrated", "frustrating", "ugh", "tired of", "fed up"),
        "stress": ("stressed", "overwhelmed", "pressure", "deadline", "can't cope", "cant cope"),
        "joy": ("happy", "great", "awesome", "excited", "love this", "thank you"),
    }
    best = ("neutral", 0.40)
    for label, words in groups.items():
        hits = sum(1 for w in words if w in t)
        if hits:
            score = min(0.90, 0.48 + 0.10 * hits)
            if score > best[1]:
                best = (label, score)
    crisis = any(
        phrase in t
        for phrase in (
            "kill myself", "end my life", "want to die", "suicide",
            "self harm", "self-harm", "hurt myself", "cut myself",
        )
    )
    return best[0], best[1], crisis


def _classify_emotion(text):
    """Run the local emotion classifier when its assets are already loaded."""
    text = (text or "").strip()
    if len(text) < 3:
        return "neutral", 0.40, False
    crisis = any(
        phrase in text.lower()
        for phrase in (
            "kill myself", "end my life", "want to die", "suicide",
            "self harm", "self-harm", "hurt myself", "cut myself",
        )
    )
    session, tokenizer = _emotion_session, _emotion_tokenizer
    if session is None or tokenizer is None:
        label, score, fallback_crisis = _emotion_fallback(text)
        return label, score, crisis or fallback_crisis
    try:
        import numpy as np
        enc = tokenizer(text, return_tensors="np", truncation=True, max_length=256, padding=True)
        input_names = {item.name for item in session.get_inputs()}
        feed = {}
        for name in input_names:
            if name in enc:
                feed[name] = enc[name]
            elif name == "token_type_ids":
                feed[name] = np.zeros_like(enc["input_ids"])
        logits = session.run(None, feed)[0][0]
        probs = 1.0 / (1.0 + np.exp(-logits))
        idx = int(np.argmax(probs))
        return EMOTION_LABELS[idx], float(probs[idx]), crisis
    except Exception:
        label, score, fallback_crisis = _emotion_fallback(text)
        return label, score, crisis or fallback_crisis


def emotion_guidance(label, score, crisis):
    """Create a safe, non-diagnostic emotional adaptation prompt for ArkaAI."""
    if crisis:
        return (
            "A possible immediate safety concern is present in the user's wording. "
            "Respond with empathy and calm. Encourage the user to contact a trusted person, "
            "local emergency services or a crisis line in their country if they may be in immediate danger. "
            "Do not provide instructions for self-harm, suicide or concealment. "
            "Do not shame, diagnose or claim certainty about their mental state. "
        )
    if label in {"sadness", "grief", "disappointment", "remorse"}:
        return (
            "The local emotion classifier suggests a low-mood signal. "
            "Acknowledge the feeling briefly, be warm and patient, and reduce cognitive load. "
            "Offer one manageable study step rather than a large plan. "
        )
    if label in {"fear", "nervousness"}:
        return (
            "The local emotion classifier suggests worry or nervousness. "
            "Use reassuring language without dismissing the concern and break the study task into small steps. "
        )
    if label in {"anger", "annoyance", "disapproval"}:
        return (
            "The local emotion classifier suggests frustration or irritation. "
            "Stay calm, avoid sounding defensive, and make the next explanation direct and practical. "
        )
    if label in {"confusion"}:
        return (
            "The local emotion classifier suggests confusion. "
            "Start from the simplest mental model, use an analogy if helpful, and check understanding before adding detail. "
        )
    if label in {"excitement", "joy", "optimism", "gratitude", "pride", "amusement"}:
        return (
            "The local emotion classifier suggests positive engagement. "
            "Match the energy while keeping the explanation useful and focused. "
        )
    return (
        "Use the user's wording and context to choose a warm, natural tone. "
    )



# Local model transport and selection are provided by model_router.py.


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, "progress.json")
PROFILE_FILE = os.path.join(BASE_DIR, "profile.json")
MISTAKES_FILE = os.path.join(BASE_DIR, "mistakes.json")
TOPICS_DIR = os.path.join(BASE_DIR, "topics")
os.makedirs(TOPICS_DIR, exist_ok=True)


def now_text():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def safe_filename(name):
    cleaned = re.sub(r"[^a-zA-Z0-9_-]+", "_", (name or "").strip())
    return cleaned.strip("._-") or "topic"


def load_json(path, default=None):
    if default is None:
        default = {}
    if not os.path.exists(path):
        return default
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError, TypeError, ValueError):
        return default


def save_json(path, data):
    os.makedirs(os.path.dirname(path) or BASE_DIR, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def topic_paths(name):
    base = os.path.join(TOPICS_DIR, safe_filename(name))
    return base + ".json", base + ".doc.txt", base + ".files.json", base + ".notes.txt"


def topic_assets_dir(name):
    path = os.path.join(TOPICS_DIR, safe_filename(name) + "_assets")
    os.makedirs(path, exist_ok=True)
    return path


def list_topics():
    names = []
    for filename in os.listdir(TOPICS_DIR):
        if filename.endswith(".json") and not filename.endswith(".files.json"):
            names.append(filename[:-5])
    return sorted(set(names), key=str.lower)


def load_topic(name):
    p = topic_paths(name)[0]
    return load_json(p, []) if os.path.exists(p) else []


def save_topic(name, messages):
    save_json(topic_paths(name)[0], messages)


def load_doc(name):
    p = topic_paths(name)[1]
    try:
        with open(p, "r", encoding="utf-8") as f:
            return f.read()
    except OSError:
        return ""


def save_doc(name, text):
    with open(topic_paths(name)[1], "w", encoding="utf-8") as f:
        f.write(text)


def load_files(name):
    p = topic_paths(name)[2]
    return load_json(p, []) if os.path.exists(p) else []


def save_files(name, items):
    save_json(topic_paths(name)[2], items)


def load_notes(name):
    p = topic_paths(name)[3]
    try:
        with open(p, "r", encoding="utf-8") as f:
            return f.read()
    except OSError:
        return ""


def save_notes(name, text):
    with open(topic_paths(name)[3], "w", encoding="utf-8") as f:
        f.write(text)


def chunk_text(text, size=900, overlap=120):
    words = re.findall(r"\S+", text or "")
    chunks = []
    start = 0
    while start < len(words):
        end = min(len(words), start + size)
        chunks.append(" ".join(words[start:end]))
        if end >= len(words):
            break
        start = max(start + 1, end - overlap)
    return chunks


def find_relevant_chunks(question, text, k=5):
    chunks = chunk_text(text)
    if not chunks:
        return []
    q = set(re.findall(r"[a-z0-9]+", (question or "").lower()))
    scored = []
    for i, chunk in enumerate(chunks):
        words = set(re.findall(r"[a-z0-9]+", chunk.lower()))
        overlap = len(q & words)
        scored.append((overlap, -i, chunk))
    scored.sort(reverse=True)
    chosen = [c for score, _, c in scored[:k] if score > 0]
    return chosen or chunks[: min(k, len(chunks))]


def extract_pdf(path):
    from PyPDF2 import PdfReader
    reader = PdfReader(path)
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        pages.append(f"[SOURCE: {os.path.basename(path)} | PDF page {i}]\n{text}")
    combined = "\n\n".join(pages).strip()
    # Some vector/scanned PDFs produce little text with PyPDF2. Try PyMuPDF as a stronger fallback.
    if len(re.sub(r"\s+", "", combined)) < 600:
        try:
            import fitz
            doc = fitz.open(path)
            alt = []
            for i, page in enumerate(doc, 1):
                text = page.get_text("text") or ""
                alt.append(f"[SOURCE: {os.path.basename(path)} | PDF page {i}]\n{text}")
            doc.close()
            alt_combined = "\n\n".join(alt).strip()
            if len(alt_combined) > len(combined):
                combined = alt_combined
        except Exception:
            pass
    return combined, f"PDF ({len(reader.pages)} pages)", []


def extract_ppt(path):
    from pptx import Presentation
    prs = Presentation(path)
    slides, embedded = [], []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
            if getattr(shape, "shape_type", None) == 13:
                try:
                    blob = shape.image.blob
                    ext = shape.image.ext or "png"
                    embedded.append((f"slide_{i}_image_{len(embedded)+1}.{ext}", blob))
                except Exception:
                    pass
        slides.append(f"[SOURCE: {os.path.basename(path)} | Slide {i}]\n" + "\n".join(texts))
    return "\n\n".join(slides).strip(), f"PowerPoint ({len(prs.slides)} slides)", embedded


def extract_image(path):
    from PIL import Image
    img = Image.open(path)
    parts = [f"[SOURCE: {os.path.basename(path)} | Image]", f"Image size: {img.width} x {img.height}"]
    try:
        import pytesseract
        ocr = pytesseract.image_to_string(img).strip()
    except Exception:
        ocr = ""
    parts.append("[OCR text]\n" + (ocr or "No OCR text available."))
    return "\n".join(parts), "Image", []


def extract_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".pptx":
        return extract_ppt(path)
    if ext == ".ppt":
        raise RuntimeError("Please save old .ppt files as .pptx before importing.")
    if ext in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"}:
        return extract_image(path)
    if ext in {".txt", ".md", ".csv"}:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f"[SOURCE: {os.path.basename(path)}]\n" + f.read(), "Text file", []
    raise RuntimeError("Supported files: PDF, PPTX, PNG, JPG/JPEG, BMP, GIF, WEBP, TXT, MD, CSV.")


def render_pdf_images(path, out_dir, max_pages=6):
    out = []
    try:
        import fitz
        doc = fitz.open(path)
        for i in range(min(max_pages, len(doc))):
            pix = doc.load_page(i).get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False)
            p = os.path.join(out_dir, f"page_{i + 1}.png")
            pix.save(p)
            out.append(p)
        doc.close()
    except Exception:
        pass
    return out


def maybe_render_ppt(path, out_dir):
    exe = None
    for candidate in ("soffice", "libreoffice"):
        try:
            subprocess.run([candidate, "--version"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
            exe = candidate
            break
        except OSError:
            pass
    if not exe:
        return []
    try:
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", out_dir, path],
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=90, check=False)
        pdf = os.path.join(out_dir, Path(path).stem + ".pdf")
        return render_pdf_images(pdf, out_dir, max_pages=8) if os.path.exists(pdf) else []
    except Exception:
        return []


def open_path(path):
    try:
        if os.name == "nt":
            os.startfile(path)  # type: ignore[attr-defined]
        elif shutil.which("xdg-open"):
            subprocess.Popen(["xdg-open", path])
        else:
            messagebox.showinfo("File", path)
    except Exception as exc:
        messagebox.showerror("Open file", str(exc))


QUESTIONS = [
    {"id": "p1", "topic": "Physics", "q": "What is the SI unit of force?", "options": ["Newton", "Joule", "Watt", "Pascal"], "answer": "Newton"},
    {"id": "p2", "topic": "Physics", "q": "Formula for velocity?", "options": ["Distance/Time", "Force*Mass", "Time/Distance", "Mass/Volume"], "answer": "Distance/Time"},
    {"id": "p3", "topic": "Physics", "q": "Unit of electric current?", "options": ["Ampere", "Volt", "Ohm", "Watt"], "answer": "Ampere"},
    {"id": "p4", "topic": "Physics", "q": "Newton's first law is also called?", "options": ["Law of Inertia", "Law of Gravity", "Law of Motion", "Law of Energy"], "answer": "Law of Inertia"},
    {"id": "c1", "topic": "Chemistry", "q": "Atomic number of Hydrogen?", "options": ["1", "2", "0", "3"], "answer": "1"},
    {"id": "c2", "topic": "Chemistry", "q": "pH of pure water?", "options": ["7", "0", "14", "5"], "answer": "7"},
]
QUESTIONS_BY_ID = {q["id"]: q for q in QUESTIONS}


class ArkaAIBaseApp:
    def __init__(self, root):
        self.root = root
        self.root.title("ArkaAI - Personal AI Study Notebook")
        self.root.geometry("1260x800")
        self.root.minsize(1040, 680)
        self.progress = load_json(DATA_FILE)
        self.mistakes = load_json(MISTAKES_FILE)
        self.profile = load_json(PROFILE_FILE)
        self.current_topic = None
        self.messages = []
        self.doc_text = ""
        self.attachments = []
        self.notes_text = ""
        self.waiting_for_answer = False
        self.current_qid = None
        self.queue = []
        self.qpos = 0
        self.round_log = []
        self.awaiting_voice = False
        self.current_emotion = "neutral"
        self.current_emotion_score = 0.0
        self.current_crisis_flag = False
        self.status_text = tk.StringVar(value="Starting local AI…")
        self._build_ui()
        self.refresh_ai_status()
        # Warm up the optional emotional-support model in the background so the first chat is smoother.
        threading.Thread(target=_emotion_model_load, daemon=True).start()
        if self.profile and list_topics():
            self.load_topic(list_topics()[0])
        else:
            self.create_topic("General")
            if not self.profile:
                self.start_onboarding()

    # ----- UI -----
    def _ui_button(self, parent, text, command, bg, fg="white", active_bg=None,
                   font=None, padx=12, pady=7, width=None):
        if active_bg is None:
            active_bg = bg
        kwargs = {
            "text": text,
            "command": command,
            "bg": bg,
            "fg": fg,
            "activebackground": active_bg,
            "activeforeground": fg,
            "relief": "flat",
            "bd": 0,
            "highlightthickness": 0,
            "cursor": "hand2",
            "font": font or ("Segoe UI", 10, "bold"),
            "padx": padx,
            "pady": pady,
        }
        if width is not None:
            kwargs["width"] = width
        btn = tk.Button(parent, **kwargs)
        btn.bind("<Enter>", lambda _e: btn.configure(bg=active_bg))
        btn.bind("<Leave>", lambda _e: btn.configure(bg=bg))
        return btn

    def _build_ui(self):
        C = {
            "bg": "#F4F6FA",
            "panel": "#FFFFFF",
            "text": "#151923",
            "muted": "#70798A",
            "border": "#DEE3EC",
            "sidebar": "#11131A",
            "sidebar2": "#1A1D26",
            "sidebar_hover": "#252A37",
            "purple": "#6C63FF",
            "purple_dark": "#554CE6",
            "teal": "#10B981",
            "blue": "#3B82F6",
            "cyan": "#06B6D4",
            "amber": "#F59E0B",
            "pink": "#EC4899",
            "indigo": "#6366F1",
            "green": "#22C55E",
            "red": "#EF4444",
        }
        self.UI_COLORS = C
        self.root.configure(bg=C["bg"])

        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure("Study.TNotebook", background=C["bg"], borderwidth=0, tabmargins=(0, 0, 0, 0))
        style.configure("Study.TNotebook.Tab", background="#E9EDF4", foreground="#4B5563",
                        padding=(16, 9), borderwidth=0, font=("Segoe UI", 9, "bold"))
        style.map("Study.TNotebook.Tab",
                  background=[("selected", C["panel"]), ("active", "#EEF0F6")],
                  foreground=[("selected", C["purple_dark"]), ("active", C["text"])])
        style.configure("Study.Treeview", rowheight=34, font=("Segoe UI", 10), borderwidth=0)
        style.configure("Study.Treeview.Heading", font=("Segoe UI", 9, "bold"))

        # Sidebar
        sidebar = tk.Frame(self.root, bg=C["sidebar"], width=286)
        sidebar.pack(side="left", fill="y")
        sidebar.pack_propagate(False)

        brand = tk.Frame(sidebar, bg=C["sidebar"])
        brand.pack(fill="x", padx=18, pady=(18, 13))
        logo = tk.Frame(brand, bg=C["purple"], width=42, height=42)
        logo.pack(side="left")
        logo.pack_propagate(False)
        tk.Label(logo, text="A", bg=C["purple"], fg="white", font=("Segoe UI", 18, "bold")).pack(expand=True)
        bt = tk.Frame(brand, bg=C["sidebar"])
        bt.pack(side="left", padx=10)
        tk.Label(bt, text="ArkaAI", bg=C["sidebar"], fg="white", font=("Segoe UI", 19, "bold")).pack(anchor="w")
        tk.Label(bt, text="Your AI study notebook", bg=C["sidebar"], fg="#9CA3AF", font=("Segoe UI", 8)).pack(anchor="w")

        self._ui_button(sidebar, "+  New lesson / chat", self.new_topic_dialog,
                        bg=C["purple"], active_bg=C["purple_dark"], pady=9).pack(fill="x", padx=14, pady=(0, 5))
        self._side_button(sidebar, "Attach source", self.attach_file)
        self._side_button(sidebar, "Check local AI", self.refresh_ai_status)

        tk.Label(sidebar, text="NOTEBOOKS", fg="#7E8492", bg=C["sidebar"], font=("Segoe UI", 8, "bold")).pack(anchor="w", padx=17, pady=(18, 7))
        frame = tk.Frame(sidebar, bg=C["sidebar"])
        frame.pack(fill="both", expand=True, padx=10)
        self.topic_list = tk.Listbox(frame, bg=C["sidebar"], fg="#E8EBF2", selectbackground="#352F72",
                                     selectforeground="white", relief="flat", borderwidth=0, highlightthickness=0,
                                     font=("Segoe UI", 10), activestyle="none", exportselection=False)
        self.topic_list.pack(side="left", fill="both", expand=True)
        sb = tk.Scrollbar(frame, command=self.topic_list.yview, bg=C["sidebar2"], troughcolor=C["sidebar"],
                          activebackground=C["purple"], relief="flat", bd=0)
        sb.pack(side="right", fill="y")
        self.topic_list.configure(yscrollcommand=sb.set)
        self.topic_list.bind("<<ListboxSelect>>", self.on_topic_select)

        foot = tk.Frame(sidebar, bg=C["sidebar"])
        foot.pack(fill="x", padx=12, pady=12)
        self._ui_button(foot, "Progress", self.show_progress_in_chat, bg=C["sidebar2"], fg="#E7EAF0",
                        active_bg=C["sidebar_hover"], font=("Segoe UI", 9, "bold")).pack(side="left", fill="x", expand=True, padx=(0,4))
        self._ui_button(foot, "Mistakes", self.start_mistakes_round, bg=C["sidebar2"], fg="#E7EAF0",
                        active_bg=C["sidebar_hover"], font=("Segoe UI", 9, "bold")).pack(side="left", fill="x", expand=True, padx=(4,0))

        # Main
        main = tk.Frame(self.root, bg=C["bg"])
        main.pack(side="left", fill="both", expand=True)
        header = tk.Frame(main, bg=C["panel"], height=76, highlightbackground=C["border"], highlightthickness=1)
        header.pack(fill="x")
        header.pack_propagate(False)

        title_box = tk.Frame(header, bg=C["panel"])
        title_box.pack(side="left", fill="y", padx=20)
        self.topic_title = tk.Label(title_box, text="General", bg=C["panel"], fg=C["text"], font=("Segoe UI", 17, "bold"))
        self.topic_title.pack(anchor="w", pady=(11, 0))
        tk.Label(title_box, text="Learn, ask, organize and create from your study material.", bg=C["panel"], fg=C["muted"], font=("Segoe UI", 8)).pack(anchor="w", pady=(0,9))

        status_box = tk.Frame(header, bg=C["panel"])
        status_box.pack(side="right", fill="y", padx=18)
        self.file_status = tk.Label(status_box, text="No sources", bg="#EFF2F6", fg="#596273",
                                    font=("Segoe UI", 8, "bold"), padx=9, pady=5)
        self.file_status.pack(side="right", padx=(8,0), pady=20)
        self.ai_status = tk.Label(status_box, textvariable=self.status_text, bg="#FFF7E8", fg="#9A6700",
                                  font=("Segoe UI", 8, "bold"), padx=9, pady=5)
        self.ai_status.pack(side="right", pady=20)
        self.mood_status = tk.Label(status_box, text="Mood: neutral", bg="#EFF2F6", fg="#596273",
                                    font=("Segoe UI", 8, "bold"), padx=9, pady=5)
        self.mood_status.pack(side="right", padx=6, pady=20)

        self.tabs = ttk.Notebook(main, style="Study.TNotebook")
        self.tabs.pack(fill="both", expand=True, padx=12, pady=10)
        self.chat_tab = tk.Frame(self.tabs, bg=C["bg"])
        self.sources_tab = tk.Frame(self.tabs, bg=C["bg"])
        self.studio_tab = tk.Frame(self.tabs, bg=C["bg"])
        self.notes_tab = tk.Frame(self.tabs, bg=C["bg"])
        self.tabs.add(self.chat_tab, text="  Chat  ")
        self.tabs.add(self.sources_tab, text="  Sources  ")
        self.tabs.add(self.studio_tab, text="  Studio  ")
        self.tabs.add(self.notes_tab, text="  Notes  ")
        self._build_chat_tab()
        self._build_sources_tab()
        self._build_studio_tab()
        self._build_notes_tab()

    def _side_button(self, parent, text, command):
        C = self.UI_COLORS
        self._ui_button(parent, text, command, bg=C["sidebar2"], fg="#E7EAF0",
                        active_bg=C["sidebar_hover"], font=("Segoe UI", 10, "bold"), pady=8).pack(fill="x", padx=14, pady=3)

    def _build_chat_tab(self):
        C = self.UI_COLORS
        self.chat_tab.grid_rowconfigure(0, weight=1)
        self.chat_tab.grid_rowconfigure(1, weight=0)
        self.chat_tab.grid_rowconfigure(2, weight=0)
        self.chat_tab.grid_columnconfigure(0, weight=1)

        area = tk.Frame(self.chat_tab, bg=C["bg"])
        area.grid(row=0, column=0, sticky="nsew", padx=2, pady=(2, 0))
        area.grid_rowconfigure(0, weight=1)
        area.grid_columnconfigure(0, weight=1)

        self.chat = tk.Text(
            area, wrap="word", state="disabled", font=("Segoe UI", 11),
            bg=C["panel"], fg=C["text"], relief="flat", bd=0,
            padx=18, pady=14,
        )
        self.chat.grid(row=0, column=0, sticky="nsew")

        sb = tk.Scrollbar(
            area, command=self.chat.yview, bg="#DCE1E9", troughcolor=C["bg"],
            activebackground=C["purple"], relief="flat", bd=0,
        )
        sb.grid(row=0, column=1, sticky="ns")
        self.chat.configure(yscrollcommand=sb.set)

        self.chat.tag_configure(
            "ai", background="#EEF1FF", foreground="#1F2840", spacing1=5,
            spacing3=11, lmargin1=10, lmargin2=10, rmargin=150,
        )
        self.chat.tag_configure(
            "user", justify="right", background="#DDF8EA", foreground="#173A29",
            spacing1=5, spacing3=11, lmargin1=150, rmargin=10,
        )
        self.chat.tag_configure(
            "system", justify="center", foreground="#7A8393",
            font=("Segoe UI", 9, "italic"), spacing3=9,
        )

        actions = tk.Frame(
            self.chat_tab, bg=C["panel"], highlightbackground=C["border"], highlightthickness=1,
        )
        actions.grid(row=1, column=0, sticky="ew", padx=2, pady=(8, 6))
        tk.Label(
            actions, text="QUICK ACTIONS", bg=C["panel"], fg=C["muted"],
            font=("Segoe UI", 8, "bold"),
        ).pack(side="left", padx=(12, 7), pady=8)

        quick = [
            ("Explain simply", lambda: self.quick_prompt("Explain the current material simply like a friendly tutor."), "#EEEAFE", "#4D43C9"),
            ("Summarize", lambda: self.run_studio("summary", True), "#EAF4FF", "#2364A8"),
            ("Study Guide", lambda: self.run_studio("studyguide", True), "#E8FBF5", "#087C68"),
            ("Quiz", self.start_quiz_from_sources, "#FFF4DC", "#A46600"),
        ]
        for label, cmd, bg, fg in quick:
            self._ui_button(
                actions, label, cmd, bg=bg, fg=fg, active_bg=bg,
                font=("Segoe UI", 9, "bold"), padx=10, pady=5,
            ).pack(side="left", padx=3, pady=4)

        row = tk.Frame(
            self.chat_tab, bg=C["panel"], highlightbackground=C["border"], highlightthickness=1,
        )
        row.grid(row=2, column=0, sticky="ew", padx=2, pady=(0, 10))
        row.grid_columnconfigure(0, weight=1)

        self.entry = tk.Entry(
            row, font=("Segoe UI", 11), bg=C["panel"], fg=C["text"],
            relief="flat", bd=0, insertbackground=C["text"],
        )
        self.entry.grid(row=0, column=0, sticky="ew", padx=(12, 5), pady=11)
        self.entry.bind("<Return>", lambda _e: self.on_send())

        self.mic_btn = self._ui_button(
            row, "Mic", self.on_mic, bg="#EEF1F6", fg="#4B5563",
            active_bg="#E2E6EF", font=("Segoe UI", 9, "bold"), width=5, pady=7,
        )
        self.mic_btn.grid(row=0, column=1, padx=(0, 5), pady=5)

        self.send_btn = self._ui_button(
            row, "Send", self.on_send, bg=C["purple"], fg="white",
            active_bg=C["purple_dark"], font=("Segoe UI", 10, "bold"), padx=15, pady=8,
        )
        self.send_btn.grid(row=0, column=2, padx=(0, 8), pady=5)

    def _build_sources_tab(self):
        C = self.UI_COLORS
        top = tk.Frame(self.sources_tab, bg=C["bg"])
        top.pack(fill="x", padx=14, pady=(14,10))
        left = tk.Frame(top, bg=C["bg"])
        left.pack(side="left", fill="x", expand=True)
        tk.Label(left, text="Sources", bg=C["bg"], fg=C["text"], font=("Segoe UI",19,"bold")).pack(anchor="w")
        tk.Label(left, text="Your notebook's PDFs, PowerPoints, images and study files.", bg=C["bg"], fg=C["muted"], font=("Segoe UI",9)).pack(anchor="w", pady=(2,0))
        controls = tk.Frame(top, bg=C["bg"])
        controls.pack(side="right")
        self._ui_button(controls, "+ Add source", self.attach_file, bg=C["purple"], fg="white", active_bg=C["purple_dark"]).pack(side="left")
        self._ui_button(controls, "Open", self.open_selected_source, bg="#EDF1F7", fg="#4B5563", active_bg="#E2E7EF", font=("Segoe UI",9,"bold")).pack(side="left", padx=(6,0))
        self._ui_button(controls, "Remove", self.remove_selected_source, bg="#FDECEC", fg=C["red"], active_bg="#F9DEDE", font=("Segoe UI",9,"bold")).pack(side="left", padx=(6,0))
        self.source_list = tk.Listbox(self.sources_tab, bg=C["panel"], fg=C["text"], selectbackground="#E9E6FF", selectforeground="#4037AE", relief="flat", bd=0, highlightbackground=C["border"], highlightthickness=1, font=("Segoe UI",10), height=6)
        self.source_list.pack(fill="x", padx=14, pady=(0,8))
        preview_header = tk.Frame(self.sources_tab, bg=C["bg"])
        preview_header.pack(fill="x", padx=14, pady=(0,5))
        tk.Label(preview_header, text="Indexed content", bg=C["bg"], fg=C["text"], font=("Segoe UI",12,"bold")).pack(side="left")
        tk.Label(preview_header, text="Used as context when you ask questions about this notebook.", bg=C["bg"], fg=C["muted"], font=("Segoe UI",8)).pack(side="left", padx=9)
        shell = tk.Frame(self.sources_tab, bg=C["panel"], highlightbackground=C["border"], highlightthickness=1)
        shell.pack(fill="both", expand=True, padx=14, pady=(0,14))
        self.source_preview = tk.Text(shell, wrap="word", font=("Segoe UI",10), bg=C["panel"], fg=C["text"], relief="flat", bd=0, padx=13, pady=11)
        ps = tk.Scrollbar(shell, command=self.source_preview.yview, bg="#DCE1E9", troughcolor=C["panel"], activebackground=C["purple"], relief="flat", bd=0)
        self.source_preview.configure(yscrollcommand=ps.set)
        self.source_preview.pack(side="left", fill="both", expand=True)
        ps.pack(side="right", fill="y")

    def _build_studio_tab(self):
        C = self.UI_COLORS
        head = tk.Frame(self.studio_tab, bg=C["bg"])
        head.pack(fill="x", padx=16, pady=(15, 8))
        tk.Label(
            head, text="Study Studio", bg=C["bg"], fg=C["text"],
            font=("Segoe UI", 19, "bold"),
        ).pack(side="left")
        tk.Label(
            head, text="Turn your sources into useful study outputs.", bg=C["bg"], fg=C["muted"],
            font=("Segoe UI", 9),
        ).pack(side="left", padx=10, pady=(6, 0))
        self.studio_status = tk.Label(
            head, text="Ready", bg=C["bg"], fg="#087443",
            font=("Segoe UI", 8, "bold"),
        )
        self.studio_status.pack(side="right", padx=5)

        grid = tk.Frame(self.studio_tab, bg=C["bg"])
        grid.pack(fill="x", padx=14)
        actions = [
            ("Summary", "summary", C["purple"], "#EEEAFE"),
            ("Study Guide", "studyguide", C["teal"], "#E8FBF5"),
            ("Mind Map", "mindmap", C["indigo"], "#ECEBFF"),
            ("Report", "report", C["blue"], "#EAF4FF"),
            ("Data Table", "table", C["cyan"], "#E8FAFD"),
            ("Flashcards", "flashcards", C["pink"], "#FDEAF4"),
            ("Quiz", "quiz", C["amber"], "#FFF4DC"),
            ("Compare Sources", "compare", "#7C3AED", "#F1EAFE"),
            ("Slide Deck", "slides", "#2563EB", "#E9F0FF"),
            ("Audio Overview", "audio", "#0F9D8A", "#E7FBF7"),
            ("Video Overview", "video", "#DB2777", "#FCEAF3"),
        ]
        for idx, (label, key, accent, tint) in enumerate(actions):
            r, c = divmod(idx, 3)
            card = tk.Frame(
                grid, bg=C["panel"], highlightbackground=C["border"], highlightthickness=1,
            )
            card.grid(row=r, column=c, sticky="nsew", padx=5, pady=5)
            grid.grid_columnconfigure(c, weight=1)
            tk.Label(
                card, text=label, bg=C["panel"], fg=C["text"],
                font=("Segoe UI", 11, "bold"), anchor="w",
            ).pack(fill="x", padx=10, pady=(10, 6))
            tk.Frame(
                card, bg=accent, height=3,
            ).pack(fill="x", padx=10)
            self._ui_button(
                card, "Create", lambda k=key: self.run_studio(k),
                bg=tint, fg=accent, active_bg=tint,
                font=("Segoe UI", 9, "bold"), padx=10, pady=5,
            ).pack(anchor="w", padx=10, pady=10)

        out_head = tk.Frame(self.studio_tab, bg=C["bg"])
        out_head.pack(fill="x", padx=14, pady=(8, 3))
        tk.Label(
            out_head, text="Output", bg=C["bg"], fg=C["text"],
            font=("Segoe UI", 11, "bold"),
        ).pack(side="left")
        self._ui_button(
            out_head, "Save output", self.save_studio_output,
            bg=C["panel"], fg=C["purple_dark"], active_bg="#EEEAFE",
            font=("Segoe UI", 9, "bold"), padx=10, pady=5,
        ).pack(side="right")

        shell = tk.Frame(
            self.studio_tab, bg=C["panel"], highlightbackground=C["border"], highlightthickness=1,
        )
        shell.pack(fill="both", expand=True, padx=14, pady=(0, 12))
        self.studio_output = tk.Text(
            shell, wrap="word", font=("Segoe UI", 10), bg=C["panel"], fg=C["text"],
            relief="flat", bd=0, highlightthickness=0, padx=14, pady=12,
            state="disabled", cursor="arrow",
        )
        ss = tk.Scrollbar(
            shell, command=self.studio_output.yview, bg="#DCE1E9", troughcolor=C["panel"],
            activebackground=C["purple"], relief="flat", bd=0,
        )
        self.studio_output.configure(yscrollcommand=ss.set)
        self.studio_output.pack(side="left", fill="both", expand=True)
        ss.pack(side="right", fill="y")

    def _build_notes_tab(self):
        C = self.UI_COLORS
        top = tk.Frame(self.notes_tab, bg=C["bg"])
        top.pack(fill="x", padx=16, pady=(15,8))
        tk.Label(top, text="Notes", bg=C["bg"], fg=C["text"], font=("Segoe UI",19,"bold")).pack(side="left")
        tk.Label(top, text="Write notes and use the local AI to organize or explain them.", bg=C["bg"], fg=C["muted"], font=("Segoe UI",9)).pack(side="left", padx=10, pady=(6,0))
        self._ui_button(top, "Save Notes", self.save_notes_ui, bg=C["purple"], fg="white", active_bg=C["purple_dark"], font=("Segoe UI",9,"bold"), padx=11, pady=6).pack(side="right")
        self._ui_button(top, "AI organize", self.organize_notes, bg="#EEEAFE", fg=C["purple_dark"], active_bg="#E5E1FF", font=("Segoe UI",9,"bold"), padx=11, pady=6).pack(side="right", padx=(0,7))
        shell = tk.Frame(self.notes_tab, bg=C["panel"], highlightbackground=C["border"], highlightthickness=1)
        shell.pack(fill="both", expand=True, padx=14, pady=(0,14))
        self.notes_editor = tk.Text(shell, wrap="word", font=("Segoe UI",11), bg=C["panel"], fg=C["text"], relief="flat", bd=0, padx=16, pady=14, insertbackground=C["text"])
        ns = tk.Scrollbar(shell, command=self.notes_editor.yview, bg="#DCE1E9", troughcolor=C["panel"], activebackground=C["purple"], relief="flat", bd=0)
        self.notes_editor.configure(yscrollcommand=ns.set)
        self.notes_editor.pack(side="left", fill="both", expand=True)
        ns.pack(side="right", fill="y")

    # ----- topic/source handling -----
    # ----- topic/source handling -----
    def refresh_topic_list(self, select_name=None):
        names = list_topics()
        self.topic_list.delete(0, "end")
        for n in names:
            self.topic_list.insert("end", n)
        if select_name in names:
            idx = names.index(select_name)
            self.topic_list.selection_clear(0, "end")
            self.topic_list.selection_set(idx)
            self.topic_list.see(idx)

    def create_topic(self, name):
        name = (name or "").strip() or "General"
        if not os.path.exists(topic_paths(name)[0]):
            save_topic(name, [])
        self.load_topic(name)

    def new_topic_dialog(self):
        name = simpledialog.askstring("New lesson / chat", "Topic or lesson name:", parent=self.root)
        if name:
            self.create_topic(name)
            self.add_message(f"New notebook created: {name}. Attach sources or just talk to me about the subject.", "ai")

    def on_topic_select(self, _=None):
        selection = self.topic_list.curselection()
        if selection:
            name = self.topic_list.get(selection[0])
            if name != self.current_topic:
                self.load_topic(name)

    def load_topic(self, name):
        self.current_topic = name
        self.messages = load_topic(name)
        self.doc_text = load_doc(name)
        self.attachments = load_files(name)
        self.notes_text = load_notes(name)
        self.waiting_for_answer = False
        self.topic_title.config(text=f"📚 {name}")
        self.redraw_chat()
        self.refresh_topic_list(name)
        self.update_source_ui()
        self.notes_editor.delete("1.0", "end")
        self.notes_editor.insert("1.0", self.notes_text)

    def redraw_chat(self):
        self.chat.configure(state="normal")
        self.chat.delete("1.0", "end")
        self.chat.configure(state="disabled")
        for m in self.messages:
            self._insert_message(m.get("text", ""), m.get("who", "ai"), persist=False)

    def _insert_message(self, text, who="ai", persist=False):
        self.chat.configure(state="normal")
        prefix = (self.profile.get("ai_name", "ArkaAI") if self.profile else "ArkaAI") + ": " if who == "ai" else ("You: " if who == "user" else "")
        tag = who if who in {"ai", "user", "system"} else "ai"
        self.chat.insert("end", prefix + text + "\n\n", tag)
        self.chat.configure(state="disabled")
        self.chat.see("end")
        if persist and self.current_topic:
            self.messages.append({"who": who, "text": text, "time": now_text()})
            save_topic(self.current_topic, self.messages)

    def add_message(self, text, who="ai"):
        self._insert_message(text, who, persist=True)

    def update_source_ui(self):
        self.source_list.delete(0, "end")
        for a in self.attachments:
            self.source_list.insert("end", f"{a.get('name', 'source')} · {a.get('type', '')}")
        self.file_status.config(text=f"{len(self.attachments)} source(s)" if self.attachments else "No sources")
        self.source_preview.delete("1.0", "end")
        self.source_preview.insert("1.0", self.doc_text[:30000])

    def attach_file(self):
        if not self.current_topic:
            return
        path = filedialog.askopenfilename(parent=self.root, title="Attach study source", filetypes=[
            ("Study files", "*.pdf *.pptx *.png *.jpg *.jpeg *.bmp *.gif *.webp *.txt *.md *.csv"),
            ("PDF", "*.pdf"), ("PowerPoint", "*.pptx"), ("Images", "*.png *.jpg *.jpeg *.bmp *.gif *.webp"), ("Text", "*.txt *.md *.csv")])
        if path:
            self.add_message(f"Reading {os.path.basename(path)}…", "system")
            threading.Thread(target=self._process_file_worker, args=(path, self.current_topic), daemon=True).start()

    def _process_file_worker(self, path, topic):
        try:
            text, kind, embedded = extract_file(path)
            asset_dir = topic_assets_dir(topic)
            dest = os.path.join(asset_dir, os.path.basename(path))
            shutil.copy2(path, dest)
            visuals = []
            ext = os.path.splitext(path)[1].lower()
            if ext in {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"}:
                visuals = [dest]
            elif ext == ".pdf":
                visuals = render_pdf_images(dest, asset_dir)
            elif ext == ".pptx":
                for name, blob in embedded:
                    p = os.path.join(asset_dir, name)
                    with open(p, "wb") as f:
                        f.write(blob)
                    visuals.append(p)
                visuals.extend(maybe_render_ppt(dest, asset_dir))
            rec = {"name": os.path.basename(path), "type": kind, "added": now_text(), "source_path": dest, "visual_paths": visuals, "chars_indexed": len(text)}
            self.root.after(0, lambda: self._file_ready(text, rec))
        except Exception as exc:
            self.root.after(0, lambda: self.add_message(f"I couldn't read that source: {exc}", "system"))

    def _file_ready(self, text, record):
        self.doc_text = (self.doc_text + "\n\n" + text).strip()[-250000:]
        self.attachments.append(record)
        save_doc(self.current_topic, self.doc_text)
        save_files(self.current_topic, self.attachments)
        self.update_source_ui()
        self.add_message(f"Added **{record['name']}**. I indexed about {record['chars_indexed']:,} characters. You can now ask about the material, figures, slides, tables or concepts.", "ai")

    def open_selected_source(self):
        sel = self.source_list.curselection()
        if not sel or sel[0] >= len(self.attachments):
            return
        p = self.attachments[sel[0]].get("source_path", "")
        if os.path.exists(p):
            open_path(p)
        else:
            messagebox.showwarning("Source", "The original source file is no longer available.")

    def remove_selected_source(self):
        sel = self.source_list.curselection()
        if not sel or sel[0] >= len(self.attachments):
            return
        item = self.attachments.pop(sel[0])
        save_files(self.current_topic, self.attachments)
        self.add_message(f"Removed source: {item.get('name', 'source')}. Reopen the notebook after adding another source if you want a fresh index.", "system")
        self.update_source_ui()

    # ----- onboarding / progress -----
    def start_onboarding(self):
        self.onboard = [("your_name", "Hey! I'm your study buddy. What's your name?"), ("ai_name", "What do you want to call me?"), ("aim", "What's your main academic goal?"), ("hobby", "What do you enjoy outside studies?")]
        self.on_idx = 0
        self.ask_onboarding()

    def ask_onboarding(self):
        self.add_message(self.onboard[self.on_idx][1], "ai")

    def handle_onboarding(self, text):
        k, _ = self.onboard[self.on_idx]
        if not hasattr(self, "onboard_answers"):
            self.onboard_answers = {}
        self.onboard_answers[k] = text
        self.on_idx += 1
        if self.on_idx >= len(self.onboard):
            self.profile = self.onboard_answers
            save_json(PROFILE_FILE, self.profile)
            self.add_message(f"Nice to meet you, {self.profile.get('your_name', 'there')}! I'm {self.profile.get('ai_name', 'ArkaAI')}. Talk naturally—I’ll adapt to your study material.", "ai")
            self.stage = None
        else:
            self.ask_onboarding()

    def show_progress_in_chat(self):
        if not self.progress:
            self.add_message("No quiz data yet.", "ai")
            return
        lines = []
        for topic, s in self.progress.items():
            total = s.get("correct", 0) + s.get("wrong", 0)
            pct = (s.get("correct", 0) / total * 100) if total else 0
            lines.append(f"{topic}: {s.get('correct', 0)} correct, {s.get('wrong', 0)} wrong ({pct:.0f}%)")
        self.add_message("Progress:\n" + "\n".join(lines), "ai")

    # ----- quiz -----
    def start_quiz_flow(self, *_):
        ids = list(QUESTIONS_BY_ID)
        random.shuffle(ids)
        self.queue = ids[:5]
        self.qpos = 0
        self.round_log = []
        self.waiting_for_answer = False
        self.current_qid = None
        self.add_message("Let's do a quick quiz. You can answer naturally.", "ai")
        self.ask_next_question()

    def start_quiz_from_sources(self):
        if not self.doc_text:
            self.start_quiz_flow()
            return
        self.run_studio("quiz", auto_chat=True)

    def start_mistakes_round(self):
        self.add_message("Mistake review is available from the original progress system. Use the built-in quiz round for saved questions.", "ai")
        self.start_quiz_flow()

    def ask_next_question(self):
        if self.qpos >= len(self.queue):
            self.finish_quiz()
            return
        q = QUESTIONS_BY_ID[self.queue[self.qpos]]
        self.current_qid = q["id"]
        self.waiting_for_answer = True
        self.add_message(f"[{q['topic']}] {q['q']}\nOptions: " + " | ".join(q["options"]), "ai")

    def finish_quiz(self):
        correct = sum(1 for r in self.round_log if r.get("correct"))
        self.add_message(f"Quiz complete — {correct}/{len(self.round_log)} correct.", "ai")
        self.waiting_for_answer = False
        self.current_qid = None

    def handle_quiz_answer(self, answer):
        q = QUESTIONS_BY_ID[self.current_qid]
        normalized = re.sub(r"\s+", "", answer.lower())
        ok = any(normalized == re.sub(r"\s+", "", opt.lower()) for opt in q["options"] if opt == q["answer"] or True) and normalized == re.sub(r"\s+", "", q["answer"].lower())
        self.round_log.append({"id": q["id"], "correct": ok})
        s = self.progress.setdefault(q["topic"], {"correct": 0, "wrong": 0})
        s["correct" if ok else "wrong"] += 1
        save_json(DATA_FILE, self.progress)
        if ok:
            self.add_message("Correct ✅", "ai")
        else:
            self.add_message(f"Not quite. The answer is **{q['answer']}**.", "ai")
        self.qpos += 1
        self.waiting_for_answer = False
        self.root.after(200, self.ask_next_question)

    # ----- AI conversation -----
    def build_context(self, user_text):
        snippets = find_relevant_chunks(user_text, self.doc_text, 6) if self.doc_text else []
        return "\n\n---\n\n".join(snippets)

    def profile_context(self):
        if not self.profile:
            return ""
        return (
            f"Student name: {self.profile.get('your_name', '')}. "
            f"Academic goal: {self.profile.get('aim', '')}."
        )

    def visual_paths_for_question(self, text):
        lower = text.lower()
        wants_visual = any(
            k in lower
            for k in (
                "image", "figure", "diagram", "graph", "chart",
                "photo", "picture", "slide", "page", "drawing",
                "table", "visual", "screenshot", "illustration",
            )
        )
        if not wants_visual:
            return []
        paths = []
        for a in self.attachments:
            paths.extend(a.get("visual_paths", []))
        return [p for p in paths if os.path.exists(p)][:4]

    def update_mood_badge(self, label, score, crisis):
        self.current_emotion = label
        self.current_emotion_score = score
        self.current_crisis_flag = crisis
        if hasattr(self, "mood_status"):
            color_map = {
                "sadness": ("#FDECEC", "#B42318"),
                "grief": ("#FDECEC", "#B42318"),
                "fear": ("#FFF4E5", "#9A6700"),
                "nervousness": ("#FFF4E5", "#9A6700"),
                "anger": ("#FDECEC", "#B42318"),
                "annoyance": ("#FFF4E5", "#9A6700"),
                "confusion": ("#EEF1FF", "#4D43C9"),
                "joy": ("#E7F8EF", "#087443"),
                "excitement": ("#E7F8EF", "#087443"),
                "optimism": ("#E7F8EF", "#087443"),
                "neutral": ("#EFF2F6", "#596273"),
            }
            bg, fg = color_map.get(
                label,
                ("#EFF2F6", "#596273"),
            )
            if crisis:
                label_text = "Support check"
                bg, fg = "#FDECEC", "#B42318"
            else:
                label_text = f"Mood: {label}"
            self.mood_status.config(
                text=label_text,
                bg=bg,
                fg=fg,
            )

    def emotional_prompt(self, label, score, crisis):
        return (
            "Local emotional-support signal: "
            f"{label} (confidence {score:.2f}). "
            "This is a probabilistic text signal, not a diagnosis. "
            + emotion_guidance(label, score, crisis)
            + "Do not announce the label as a medical fact. "
            "Keep the user in control and focus on useful study support."
        )

    def quick_prompt(self, prompt):
        if getattr(self, "stream_active", False):
            return
        self.add_message(prompt, "user")
        self.handle_general_chat(prompt)

    def _begin_stream_message(self):
        self.stream_active = True
        self.stream_text = ""
        self.chat.configure(state="normal")
        self.chat.insert("end", "ArkaAI: ", "ai")
        self.chat.configure(state="disabled")
        self.chat.see("end")

    def _append_stream_chunk(self, chunk):
        if not getattr(self, "stream_active", False):
            return
        self.stream_text += chunk
        self.chat.configure(state="normal")
        self.chat.insert("end", chunk, "ai")
        self.chat.configure(state="disabled")
        self.chat.see("end")

    def _finish_stream_message(self, answer, error):
        self.stream_active = False
        self.send_btn.config(state="normal")
        self.mic_btn.config(state="normal")
        if error:
            self.chat.configure(state="normal")
            self.chat.insert("end", "\nAI error: " + error + "\n\n", "system")
            self.chat.configure(state="disabled")
            self.chat.see("end")
            return
        final = (answer or self.stream_text).strip()
        if final and self.current_topic:
            self.messages.append(
                {
                    "who": "ai",
                    "text": final,
                    "time": now_text(),
                }
            )
            save_topic(
                self.current_topic,
                self.messages,
            )
        self.chat.configure(state="normal")
        self.chat.insert("end", "\n\n", "ai")
        self.chat.configure(state="disabled")
        self.chat.see("end")

    def on_send(self):
        text = self.entry.get().strip()
        if not text:
            return
        self.entry.delete(0, "end")
        self.add_message(text, "user")
        self.route_message(text)

    def route_message(self, text):
        lower = text.lower().strip()
        if hasattr(self, "onboard") and getattr(self, "on_idx", 0) < len(getattr(self, "onboard", [])) and not self.profile:
            self.handle_onboarding(text)
            return
        if self.waiting_for_answer and self.current_qid:
            self.handle_quiz_answer(text)
            return
        if lower in {"progress", "show progress", "my progress"}:
            self.show_progress_in_chat(); return
        if lower in {"quiz", "quiz me", "start quiz"}:
            self.start_quiz_from_sources(); return
        if lower in {"new round", "new quiz"}:
            self.start_quiz_flow(); return
        if lower in {"summarize", "summary", "summarise"}:
            self.run_studio("summary", auto_chat=True); return
        self.handle_general_chat(text)

    # ----- Studio -----
    def require_sources(self):
        if not self.doc_text:
            self.add_message(
                "Attach at least one study source first. Then Studio can create grounded material from it.",
                "system",
            )
            return False
        return True

    def source_prompt(self):
        return self.doc_text[:180000]

    def studio_call(self, instruction, context=None):
        ctx = context if context is not None else self.source_prompt()
        messages = [
            {
                "role": "system",
                "content": (
                    "You are the ArkaAI Study Studio engine. Work primarily from supplied source material. "
                    "Preserve source terminology, cite source markers for source-derived claims, and do not invent missing facts."
                ),
            },
            {
                "role": "user",
                "content": instruction + "\n\nSOURCE MATERIAL:\n" + ctx,
            },
        ]
        return backend_chat(
            messages,
            MODEL_NAME,
            temperature=0.25,
            num_ctx=8192,
            timeout=300,
        )

    def _set_studio_text(self, text, append=False):
        self.studio_output.configure(state="normal")
        if not append:
            self.studio_output.delete("1.0", "end")
        self.studio_output.insert("end", text)
        self.studio_output.configure(state="disabled")
        self.studio_output.see("end")

    def run_studio(self, kind, auto_chat=False):
        if not self.require_sources():
            return

        prompts = {
            "summary": "Create a clear study summary with headings, key ideas, definitions and source markers.",
            "studyguide": "Create an exam-ready study guide with objectives, core concepts, definitions, formulas, examples, common mistakes and last-minute revision.",
            "mindmap": "Create a clean text mind map using an indented tree from central topic to major branches and sub-concepts.",
            "report": "Write a structured study report with title, introduction, key sections, evidence from the sources, conclusion and source markers.",
            "table": "Extract useful facts, comparisons and relationships into Markdown tables supported directly by the sources.",
            "flashcards": "Create 20 high-value flashcards. Use FRONT: and BACK: labels and cover important concepts from the sources.",
            "quiz": "Create 10 study questions from the sources with answers. Mix multiple-choice and short-answer questions and label answers clearly.",
            "compare": "Compare main ideas across supplied sources. Identify agreements, differences, complementary points and gaps with source markers.",
            "slides": "Create a 10-slide presentation outline. For each slide provide TITLE and 3-6 concise bullets faithful to the sources.",
            "audio": "Write a natural 6-8 minute Audio Overview script in a friendly two-person tutor discussion style grounded in the sources.",
            "video": "Create an 8-scene Video Overview storyboard. For each scene provide VISUAL, ON-SCREEN TEXT, NARRATION and SOURCE. This is a storyboard/script, not a claim of a rendered video.",
        }

        instruction = prompts.get(kind)
        if not instruction:
            return

        self.tabs.select(self.studio_tab)
        self.studio_status.config(text="Generating…", fg="#9A6700")
        self._set_studio_text("Generating…")

        def worker():
            try:
                messages = [
                    {
                        "role": "system",
                        "content": (
                            "You are the ArkaAI Study Studio engine. "
                            "Work primarily from the supplied study sources, preserve terminology, cite source markers and never invent source facts."
                        ),
                    },
                    {
                        "role": "user",
                        "content": instruction + "\n\nSOURCE MATERIAL:\n" + self.source_prompt(),
                    },
                ]

                pieces = []
                first = True
                for chunk in backend_stream(
                    messages,
                    _pick_model(),
                    temperature=0.25,
                    num_ctx=8192,
                    timeout=300,
                ):
                    pieces.append(chunk)
                    if first:
                        self.root.after(0, lambda: self._set_studio_text(""))
                        first = False
                    self.root.after(
                        0,
                        lambda c=chunk: self._set_studio_text(c, append=True),
                    )

                result = "".join(pieces).strip()

                if kind == "slides":
                    result += "\n\n[Tip] Save output and choose to create a PowerPoint deck."
                elif kind == "audio":
                    artifact = self._create_audio_overview(result)
                    result += (
                        f"\n\n[Audio file] {artifact}"
                        if artifact
                        else "\n\n[Audio note] Script generated; local audio rendering was unavailable."
                    )
                elif kind == "video":
                    artifact = self._create_video_overview(result)
                    result += (
                        f"\n\n[Video file] {artifact}"
                        if artifact
                        else "\n\n[Video note] Storyboard generated; install FFmpeg for MP4 rendering."
                    )

                self.root.after(
                    0,
                    lambda r=result: self._studio_done(r, auto_chat, False),
                )

            except Exception as exc:
                self.root.after(
                    0,
                    lambda e=str(exc): self._studio_done(
                        "Studio error: " + e,
                        auto_chat,
                        True,
                    ),
                )

        threading.Thread(
            target=worker,
            daemon=True,
        ).start()

    def _studio_done(self, result, auto_chat, error=False):
        self._set_studio_text(result)
        self.studio_status.config(
            text="Error" if error else "Ready",
            fg="#B42318" if error else "#087443",
        )
        if auto_chat and not error:
            self.add_message(result, "ai")

    def _create_audio_overview(self, script):
        try:
            import pyttsx3
            out = os.path.join(topic_assets_dir(self.current_topic), f"Audio_Overview_{datetime.now().strftime('%Y%m%d_%H%M%S')}.wav")
            engine = pyttsx3.init()
            engine.setProperty("rate", 165)
            engine.save_to_file(script[:12000], out)
            engine.runAndWait()
            return out if os.path.exists(out) else None
        except Exception as exc:
            print("Audio overview rendering skipped:", exc)
            return None

    def _create_video_overview(self, storyboard):
        # Create a simple educational MP4 from storyboard scenes when FFmpeg is available.
        # Without FFmpeg, the scene PNGs are still saved so the artifact is not lost.
        try:
            from PIL import Image, ImageDraw, ImageFont
            asset_dir = topic_assets_dir(self.current_topic)
            video_dir = os.path.join(asset_dir, "video_overview")
            os.makedirs(video_dir, exist_ok=True)
            scenes = re.split(r"\n\s*(?=SCENE\s*\d+)|\n\s*(?=Scene\s*\d+)", storyboard, flags=re.I)
            scenes = [x.strip() for x in scenes if x.strip()]
            if not scenes:
                scenes = [storyboard]
            try:
                font = ImageFont.truetype("arial.ttf", 30)
                small = ImageFont.truetype("arial.ttf", 22)
            except Exception:
                font = ImageFont.load_default()
                small = ImageFont.load_default()
            image_paths = []
            for i, scene in enumerate(scenes[:12], 1):
                img = Image.new("RGB", (1280, 720), "white")
                draw = ImageDraw.Draw(img)
                title = f"ArkaAI Video Overview — Scene {i}"
                draw.text((55, 45), title, fill="black", font=font)
                lines = textwrap.wrap(scene, width=70)
                y = 120
                for line in lines[:22]:
                    draw.text((60, y), line, fill="black", font=small)
                    y += 27
                out = os.path.join(video_dir, f"scene_{i:02d}.png")
                img.save(out)
                image_paths.append(out)
            ffmpeg = None
            for candidate in ("ffmpeg", "ffmpeg.exe"):
                if shutil.which(candidate):
                    ffmpeg = candidate
                    break
            if not ffmpeg:
                return None
            mp4 = os.path.join(asset_dir, f"Video_Overview_{datetime.now().strftime('%Y%m%d_%H%M%S')}.mp4")
            # Use the generated scene PNGs directly. Each scene stays on screen for 4 seconds.
            subprocess.run([ffmpeg, "-y", "-framerate", "1/4", "-i", os.path.join(video_dir, "scene_%02d.png"),
                            "-c:v", "libx264", "-pix_fmt", "yuv420p", mp4],
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180, check=False)
            return mp4 if os.path.exists(mp4) else None
        except Exception as exc:
            print("Video overview rendering skipped:", exc)
            return None

    def save_studio_output(self):
        text = self.studio_output.get("1.0", "end").strip()
        if not text or text == "Generating…":
            return
        out_dir = topic_assets_dir(self.current_topic)
        path = os.path.join(out_dir, f"studio_{datetime.now().strftime('%Y%m%d_%H%M%S')}.md")
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
        messagebox.showinfo("Saved", f"Saved Studio output to:\n{path}")

        # Also offer real PPTX creation for slide outlines.
        if "TITLE:" in text and "slide" in text.lower():
            self.maybe_create_pptx(text)

    def maybe_create_pptx(self, outline):
        answer = messagebox.askyesno("Slide deck", "Create a PowerPoint (.pptx) from this slide outline too?", parent=self.root)
        if not answer:
            return
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            prs = Presentation()
            blocks = re.split(r"\n(?=\s*TITLE\s*:)", outline, flags=re.I)
            made = 0
            for block in blocks:
                mt = re.search(r"TITLE\s*:\s*(.+)", block, re.I)
                if not mt:
                    continue
                title = mt.group(1).strip()
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = title
                tf = slide.placeholders[1].text_frame
                tf.clear()
                bullets = re.findall(r"^\s*[-•]\s+(.+)$", block, re.M)
                if not bullets:
                    bullets = [x.strip() for x in block.splitlines() if x.strip() and not re.match(r"\s*TITLE\s*:", x, re.I)][:5]
                for i, bullet in enumerate(bullets[:7]):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = bullet
                    p.level = 0
                    p.font.size = Pt(18)
                made += 1
            if made == 0:
                raise RuntimeError("No slide titles were detected in the outline.")
            out = os.path.join(topic_assets_dir(self.current_topic), f"ArkaAI_Slides_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx")
            prs.save(out)
            messagebox.showinfo("Slide deck created", out)
        except Exception as exc:
            messagebox.showerror("PowerPoint", str(exc))

    # ----- Notes -----
    def save_notes_ui(self):
        self.notes_text = self.notes_editor.get("1.0", "end").rstrip()
        save_notes(self.current_topic, self.notes_text)
        messagebox.showinfo("Notes", "Notes saved.")

    def organize_notes(self):
        notes = self.notes_editor.get("1.0", "end").strip()
        if not notes:
            return
        def worker():
            try:
                result = self.studio_call("Organize these student notes into clear headings, bullet points, definitions, formulas, and action items. Keep the original meaning.\n\nSTUDENT NOTES:\n" + notes, context="")
                self.root.after(0, lambda: self._replace_notes(result))
            except Exception as exc:
                self.root.after(0, lambda: messagebox.showerror("AI notes", str(exc)))
        threading.Thread(target=worker, daemon=True).start()

    def _replace_notes(self, text):
        self.notes_editor.delete("1.0", "end")
        self.notes_editor.insert("1.0", text)

    # ----- voice -----
    def on_mic(self):
        if self.awaiting_voice:
            return
        self.awaiting_voice = True
        self.mic_btn.config(state="disabled", text="🔴")
        def worker():
            try:
                import sounddevice as sd
                import scipy.io.wavfile as wavfile
                import speech_recognition as sr
                fs = 16000
                recording = sd.rec(5 * fs, samplerate=fs, channels=1, dtype="int16")
                sd.wait()
                p = os.path.join(tempfile.gettempdir(), "arkaai_voice.wav")
                wavfile.write(p, fs, recording)
                r = sr.Recognizer()
                with sr.AudioFile(p) as src:
                    audio = r.record(src)
                text = r.recognize_google(audio)
                self.root.after(0, lambda: self._voice_done(text, None))
            except Exception as exc:
                self.root.after(0, lambda: self._voice_done(None, str(exc)))
        threading.Thread(target=worker, daemon=True).start()

    def _voice_done(self, text, error):
        self.awaiting_voice = False
        self.mic_btn.config(state="normal", text="🎤")
        if error:
            self.add_message("Voice input failed: " + error, "system")
            return
        self.entry.delete(0, "end")
        self.entry.insert(0, text)
        self.on_send()

    def refresh_ai_status(self):
        server, model, msg = backend_status()
        self.status_text.set(msg)
        if model:
            self.ai_status.config(bg="#E7F8EF", fg="#087443")
        else:
            self.ai_status.config(bg="#FDECEC", fg="#B42318")
        if hasattr(self, "mood_status"):
            self.mood_status.config(
                text=f"Mood: {self.current_emotion or 'neutral'}" if not self.current_crisis_flag else "Support check",
            )


# -----------------------------------------------------------------------------
# ArkaAI intelligence services
# -----------------------------------------------------------------------------
from model_router import AdaptiveModelRouter, installed_models
from snapdragon_backend import create_onnx_session, status_text as snapdragon_status_text

ROUTER = AdaptiveModelRouter()
FAST_MODEL = ROUTER.fast_model
DEEP_MODEL = ROUTER.deep_model
VISION_MODEL = ROUTER.primary_model
MODEL_NAME = FAST_MODEL
FALLBACK_MODEL = DEEP_MODEL
EMBED_MODEL = ROUTER.embed_model

MEMORY_FILE = os.path.join(BASE_DIR, "semantic_memory.json")
LEARNING_PROFILE_FILE = os.path.join(BASE_DIR, "learning_profile.json")
SEMANTIC_DIR = os.path.join(BASE_DIR, "semantic_index")
os.makedirs(SEMANTIC_DIR, exist_ok=True)

DEFAULT_LEARNING_PROFILE = {
    "strengths": [],
    "weaknesses": [],
    "topic_stats": {},
    "goals": [],
    "preferences": {},
    "recent_topics": [],
    "confidence": {},
    "updated": "",
}


def _merge_profile(base):
    out = dict(DEFAULT_LEARNING_PROFILE)
    if isinstance(base, dict):
        out.update(base)
    for key in ("strengths", "weaknesses", "goals", "recent_topics"):
        if not isinstance(out.get(key), list):
            out[key] = []
    for key in ("topic_stats", "preferences", "confidence"):
        if not isinstance(out.get(key), dict):
            out[key] = {}
    return out


def load_learning_profile():
    return _merge_profile(load_json(LEARNING_PROFILE_FILE, {}))


def save_learning_profile(profile):
    profile["updated"] = now_text()
    save_json(LEARNING_PROFILE_FILE, profile)


def _semantic_path(topic):
    return os.path.join(SEMANTIC_DIR, safe_filename(topic) + ".json")


def _load_semantic_index(topic):
    return load_json(_semantic_path(topic), {"text_hash": "", "chunks": []})


def _save_semantic_index(topic, payload):
    save_json(_semantic_path(topic), payload)


def _text_hash(text):
    return hashlib.sha1((text or "").encode("utf-8", errors="ignore")).hexdigest()


def _cosine(a, b):
    if not a or not b:
        return -1.0
    n = min(len(a), len(b))
    dot = sum(float(a[i]) * float(b[i]) for i in range(n))
    na = math.sqrt(sum(float(x) * float(x) for x in a))
    nb = math.sqrt(sum(float(x) * float(x) for x in b))
    return dot / (na * nb) if na > 0 and nb > 0 else -1.0


def _embed_texts(texts, timeout=120):
    try:
        return ROUTER.embed(texts, timeout=timeout)
    except Exception as exc:
        print("Semantic embedding unavailable:", exc)
        return None


def _keyword_score(query, text):
    q = set(re.findall(r"[a-z0-9]+", (query or "").lower()))
    if not q:
        return 0.0
    words = set(re.findall(r"[a-z0-9]+", (text or "").lower()))
    return len(q & words) / max(1, len(q))


def _semantic_search_records(query, records, k=5):
    if not records:
        return []
    qvecs = _embed_texts([query], timeout=45)
    qvec = qvecs[0] if qvecs else None
    scored = []
    for index, record in enumerate(records):
        text = record.get("text", "")
        semantic = _cosine(qvec, record.get("embedding")) if qvec else -1.0
        lexical = _keyword_score(query, text)
        score = (0.82 * semantic + 0.18 * lexical) if semantic >= 0 else lexical
        scored.append((score, -index, record))
    scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return [record for score, _, record in scored[:k] if score > 0.0]


def rebuild_semantic_index(topic, text):
    chunks = chunk_text(text or "", size=420, overlap=70)
    records = [
        {"id": str(i + 1), "topic": topic, "text": chunk, "embedding": None}
        for i, chunk in enumerate(chunks)
    ]
    for start in range(0, len(records), 16):
        batch = records[start:start + 16]
        vectors = _embed_texts([item["text"] for item in batch], timeout=120)
        if not vectors:
            break
        for record, vector in zip(batch, vectors):
            record["embedding"] = vector
    _save_semantic_index(topic, {
        "text_hash": _text_hash(text),
        "chunks": records,
        "updated": now_text(),
    })
    return len(records), sum(1 for record in records if record.get("embedding"))


def _semantic_document_context(topic, query, k=6):
    index = _load_semantic_index(topic)
    text = load_doc(topic)
    if not text:
        return ""
    if not index.get("chunks") or index.get("text_hash") != _text_hash(text):
        fallback = find_relevant_chunks(query, text, k=k)
        threading.Thread(target=rebuild_semantic_index, args=(topic, text), daemon=True).start()
        return "\n\n---\n\n".join(fallback)
    hits = _semantic_search_records(query, index.get("chunks", []), k=k)
    if not hits:
        hits = [{"text": chunk} for chunk in find_relevant_chunks(query, text, k=k)]
    return "\n\n---\n\n".join(record.get("text", "")[:2200] for record in hits if record.get("text"))


def _load_memory_records():
    data = load_json(MEMORY_FILE, [])
    return data if isinstance(data, list) else []


def _save_memory_records(records):
    save_json(MEMORY_FILE, records[-350:])


def _remember_exchange(topic, user_text, answer):
    user_text = (user_text or "").strip()
    answer = (answer or "").strip()
    if not user_text or not answer:
        return
    record = {
        "id": f"m-{int(time.time() * 1000)}",
        "topic": topic or "General",
        "kind": "conversation",
        "text": f"User: {user_text}\nArkaAI: {answer[:1800]}",
        "created": now_text(),
        "embedding": None,
    }
    records = _load_memory_records()
    records.append(record)
    vectors = _embed_texts([record["text"]], timeout=60)
    if vectors:
        record["embedding"] = vectors[0]
    _save_memory_records(records)


def _semantic_memory_context(query, topic=None, k=5):
    records = _load_memory_records()
    if topic:
        topical = [record for record in records if record.get("topic") == topic]
        other = [record for record in records if record.get("topic") != topic]
        records = topical + other[:120]
    hits = _semantic_search_records(query, records, k=k)
    return "\n\n---\n\n".join(
        f"[{record.get('topic', 'General')} | {record.get('created', '')}] {record.get('text', '')[:1400]}"
        for record in hits
    )


def detect_emotion(text):
    """Non-blocking emotion API: fallback immediately until the local model is warm."""
    if _emotion_session is not None and _emotion_tokenizer is not None:
        return _classify_emotion(text)
    label, score, crisis = _emotion_fallback(text)
    return label, score, crisis


def backend_status():
    server, model_ok, detail, _meta = ROUTER.status()
    return server, model_ok, detail


def _pick_model(preference="fast", multimodal=False):
    return ROUTER.pick("deep" if preference == "deep" else "fast", multimodal)


def backend_chat(messages, model=None, image_paths=None, temperature=0.3, num_ctx=8192, timeout=900, num_predict=None):
    mode = "deep" if model == DEEP_MODEL and DEEP_MODEL != FAST_MODEL else "fast"
    return ROUTER.chat(messages, mode=mode, temperature=temperature, num_ctx=num_ctx,
                       num_predict=num_predict, image_paths=image_paths, timeout=timeout)


def backend_stream(messages, model=None, image_paths=None, temperature=0.3, num_ctx=8192, timeout=900, num_predict=None):
    mode = "deep" if model == DEEP_MODEL and DEEP_MODEL != FAST_MODEL else "fast"
    yield from ROUTER.stream(messages, mode=mode, temperature=temperature, num_ctx=num_ctx,
                             num_predict=num_predict, image_paths=image_paths, timeout=timeout)


class ArkaAIApp(ArkaAIBaseApp):
    """Production composition of the desktop UI and the adaptive intelligence layer."""

    def __init__(self, root):
        self.learning_profile = load_learning_profile()
        self.stream_active = False
        self.stream_text = ""
        super().__init__(root)
        self.learning_profile = load_learning_profile()

    def _build_ui(self):
        super()._build_ui()
        sidebar = self.root.winfo_children()[0]
        button = self._ui_button(
            sidebar, "Memory / Profile", self.show_learning_profile,
            bg=self.UI_COLORS["sidebar2"], fg="#E7EAF0",
            active_bg=self.UI_COLORS["sidebar_hover"],
            font=("Segoe UI", 10, "bold"), pady=8,
        )
        children = sidebar.winfo_children()
        before = children[-1] if children else None
        if before is not None:
            button.pack(fill="x", padx=14, pady=3, before=before)
        else:
            button.pack(fill="x", padx=14, pady=3)

    def load_topic(self, name):
        super().load_topic(name)
        if self.doc_text:
            index = _load_semantic_index(name)
            if index.get("text_hash") != _text_hash(self.doc_text):
                threading.Thread(target=rebuild_semantic_index,
                                 args=(name, self.doc_text), daemon=True).start()

    def _file_ready(self, text, record):
        super()._file_ready(text, record)
        if self.current_topic and self.doc_text:
            def worker():
                try:
                    chunks, vectors = rebuild_semantic_index(self.current_topic, self.doc_text)
                    message = f"Semantic search ready: {chunks} passages indexed"
                    message += (
                        f" with {vectors} local embeddings."
                        if vectors else
                        "; keyword fallback is active until embeddings are available."
                    )
                    self.root.after(0, lambda m=message: self.add_message(m, "system"))
                except Exception as exc:
                    print("Semantic index error:", exc)
            threading.Thread(target=worker, daemon=True).start()

    def build_context(self, user_text):
        parts = []
        if self.current_topic and self.doc_text:
            document = _semantic_document_context(self.current_topic, user_text, k=6)
            if document:
                parts.append("SEMANTICALLY RETRIEVED STUDY MATERIAL:\n" + document)
        memory = _semantic_memory_context(user_text, topic=self.current_topic, k=5)
        if memory:
            parts.append("SEMANTIC MEMORY FROM EARLIER STUDY SESSIONS:\n" + memory)
        return "\n\n=====\n\n".join(parts)

    def profile_context(self):
        profile = self.profile or {}
        learning = _merge_profile(getattr(self, "learning_profile", {}))
        strengths = learning.get("strengths", [])[-6:]
        weaknesses = learning.get("weaknesses", [])[-6:]
        stats = learning.get("topic_stats", {})
        recent = learning.get("recent_topics", [])[-6:]
        lines = [
            f"Student name: {profile.get('your_name', '')}",
            f"Academic goal: {profile.get('aim', '')}",
            f"Preferred assistant name: {profile.get('ai_name', 'ArkaAI')}",
        ]
        if strengths:
            lines.append("Known strengths: " + "; ".join(strengths))
        if weaknesses:
            lines.append("Known weak areas: " + "; ".join(weaknesses))
        if stats:
            compact = []
            for topic, values in list(stats.items())[-8:]:
                total = values.get("correct", 0) + values.get("wrong", 0)
                pct = round(100 * values.get("correct", 0) / total) if total else 0
                compact.append(f"{topic} ({pct}% quiz accuracy)")
            lines.append("Learning stats: " + ", ".join(compact))
        if recent:
            lines.append("Recently studied: " + ", ".join(recent))
        return "\n".join(lines)

    def _record_student_result(self, question, answer, correct):
        self.learning_profile = _merge_profile(getattr(self, "learning_profile", {}))
        topic = question.get("topic", self.current_topic or "General")
        stats = self.learning_profile["topic_stats"].setdefault(topic, {"correct": 0, "wrong": 0})
        stats["correct" if correct else "wrong"] += 1
        self.learning_profile["recent_topics"] = [
            item for item in self.learning_profile.get("recent_topics", []) if item != topic
        ] + [topic]
        if correct:
            item = f"{topic}: correctly answered — {question.get('q', '')[:130]}"
            values = [x for x in self.learning_profile["strengths"] if x != item]
            self.learning_profile["strengths"] = (values[-19:] + [item])
            self.learning_profile["confidence"][topic] = min(
                1.0, self.learning_profile["confidence"].get(topic, 0.50) + 0.08
            )
        else:
            item = (
                f"{topic}: review needed — {question.get('q', '')[:130]} "
                f"(answered: {answer[:90]})"
            )
            values = [x for x in self.learning_profile["weaknesses"] if x != item]
            self.learning_profile["weaknesses"] = (values[-19:] + [item])
            self.learning_profile["confidence"][topic] = max(
                0.0, self.learning_profile["confidence"].get(topic, 0.50) - 0.10
            )
        save_learning_profile(self.learning_profile)

    def handle_quiz_answer(self, answer):
        question = QUESTIONS_BY_ID[self.current_qid]
        normalized = re.sub(r"\s+", "", answer.lower())
        expected = re.sub(r"\s+", "", question["answer"].lower())
        correct = normalized == expected
        self.round_log.append({"id": question["id"], "correct": correct, "topic": question["topic"]})
        stats = self.progress.setdefault(question["topic"], {"correct": 0, "wrong": 0})
        stats["correct" if correct else "wrong"] += 1
        if correct:
            self.mistakes.pop(question["id"], None)
            self.add_message("Correct — I’ve recorded that as a strength.", "ai")
        else:
            self.mistakes[question["id"]] = self.mistakes.get(question["id"], 0) + 1
            self.add_message(
                f"Not quite. The answer is **{question['answer']}**. "
                "I’ll remember this weak point for future review.", "ai"
            )
        save_json(DATA_FILE, self.progress)
        save_json(MISTAKES_FILE, self.mistakes)
        self._record_student_result(question, answer, correct)
        threading.Thread(
            target=_remember_exchange,
            args=(self.current_topic,
                  f"Quiz: {question['q']} | My answer: {answer}",
                  f"Correct: {correct}. Expected answer: {question['answer']}"),
            daemon=True,
        ).start()
        self.qpos += 1
        self.waiting_for_answer = False
        self.root.after(220, self.ask_next_question)

    def show_learning_profile(self):
        learning = _merge_profile(getattr(self, "learning_profile", {}))
        lines = [
            "LEARNING PROFILE",
            "",
            f"Goal: {(self.profile or {}).get('aim', 'Not set')}",
        ]
        stats = learning.get("topic_stats", {})
        if stats:
            lines.append("\nTOPIC PROGRESS")
            for topic, values in stats.items():
                total = values.get("correct", 0) + values.get("wrong", 0)
                pct = round(100 * values.get("correct", 0) / total) if total else 0
                lines.append(
                    f"• {topic}: {values.get('correct', 0)} correct, "
                    f"{values.get('wrong', 0)} wrong ({pct}%)"
                )
        if learning.get("strengths"):
            lines.append("\nSTRENGTH MEMORY")
            lines.extend("• " + x for x in learning["strengths"][-6:])
        if learning.get("weaknesses"):
            lines.append("\nMISTAKE / WEAKNESS MEMORY")
            lines.extend("• " + x for x in learning["weaknesses"][-6:])
        lines.append(
            "\nAI ROUTING\n"
            f"• Fast: {ROUTER.fast_model}\n"
            f"• Deep: {ROUTER.deep_model}\n"
            f"• Embeddings: {ROUTER.embed_model}\n"
            f"• Platform: {snapdragon_status_text()}"
        )
        self.add_message("\n".join(lines), "ai")

    def start_mistakes_round(self):
        weak = _merge_profile(getattr(self, "learning_profile", {})).get("weaknesses", [])
        if weak:
            self.add_message(
                "Here are the weak points I’m remembering:\n" +
                "\n".join("• " + item for item in weak[-6:]), "ai"
            )
        else:
            self.add_message(
                "No saved mistake memory yet. I’ll learn your weak areas as you answer quizzes.", "ai"
            )
        self.start_quiz_flow()

    def handle_general_chat(self, text):
        self.send_btn.config(state="disabled")
        self.mic_btn.config(state="disabled")
        self._begin_stream_message()

        def worker():
            try:
                _, model_ok, status = backend_status()
                if not model_ok:
                    raise RuntimeError(status)
                label, score, crisis = detect_emotion(text)
                self.root.after(0, lambda: self.update_mood_badge(label, score, crisis))
                context = self.build_context(text)
                visuals = self.visual_paths_for_question(text)
                deep_words = (
                    "deeply", "detailed", "derive", "prove", "compare",
                    "analyze", "analyse", "reason", "solve", "debug", "why exactly"
                )
                deep = bool(context) or len(text.split()) > 28 or any(w in text.lower() for w in deep_words)
                mode = "deep" if deep else "fast"
                if visuals and deep:
                    mode = "deep"

                system = (
                    "You are ArkaAI, a local personal AI study assistant and emotional companion. "
                    "Be warm, natural, honest and highly capable. Answer directly, then add depth when useful. "
                    "Use the student's profile, semantic memory and supplied study sources when relevant. "
                    "Never invent claims from source material. Preserve source markers when present. "
                    "Do not claim to be ChatGPT or Perplexity; you are ArkaAI.\n"
                    "Run locally by default. On supported Snapdragon systems, Qualcomm QNN/NPU acceleration is used for compatible local ONNX services when available.\n\n"
                    "STUDENT PROFILE:\n" + self.profile_context() + "\n\n" +
                    self.emotional_prompt(label, score, crisis)
                )
                if context:
                    system += "\n\nRETRIEVED CONTEXT:\n" + context[:15000]

                history = [{"role": "system", "content": system}]
                recent = self.messages[-14:]
                for message in recent:
                    if message.get("who") not in ("user", "ai"):
                        continue
                    if (message is recent[-1] and message.get("who") == "user" and
                            message.get("text") == text):
                        continue
                    history.append({
                        "role": "user" if message.get("who") == "user" else "assistant",
                        "content": message.get("text", ""),
                    })
                history.append({"role": "user", "content": text})

                full = []
                for chunk in ROUTER.stream(
                    history,
                    mode=mode,
                    temperature=0.22 if mode == "deep" else 0.30,
                    num_ctx=16384 if mode == "deep" else 8192,
                    num_predict=850 if mode == "deep" else 320,
                    image_paths=visuals,
                    timeout=1200,
                ):
                    full.append(chunk)
                    self.root.after(0, lambda c=chunk: self._append_stream_chunk(c))

                answer = "".join(full).strip()
                if answer:
                    threading.Thread(
                        target=_remember_exchange,
                        args=(self.current_topic, text, answer),
                        daemon=True,
                    ).start()
                    learning = _merge_profile(getattr(self, "learning_profile", {}))
                    topic = self.current_topic or "General"
                    recent_topics = [x for x in learning.get("recent_topics", []) if x != topic]
                    learning["recent_topics"] = (recent_topics + [topic])[-10:]
                    save_learning_profile(learning)
                    self.learning_profile = learning
                self.root.after(0, lambda a=answer: self._finish_stream_message(a, None))
            except Exception as exc:
                self.root.after(0, lambda e=str(exc): self._finish_stream_message("", e))

        threading.Thread(target=worker, daemon=True).start()

    def studio_call(self, instruction, context=None):
        source = context if context is not None else self.source_prompt()
        messages = [
            {
                "role": "system",
                "content": (
                    "You are the ArkaAI Study Studio engine. Work primarily from supplied source material. "
                    "Preserve terminology, cite source markers, do not invent missing facts, and produce clear student-friendly output.\n\n"
                    + self.profile_context()
                ),
            },
            {"role": "user", "content": instruction + "\n\nSOURCE MATERIAL:\n" + source[:220000]},
        ]
        return ROUTER.chat(
            messages, mode="deep", temperature=0.18, num_ctx=16384,
            num_predict=1400, timeout=1800,
        )

    def refresh_ai_status(self):
        _server, model_ok, detail = backend_status()
        self.status_text.set(detail + " · " + snapdragon_status_text())
        if model_ok:
            self.ai_status.config(bg="#E7F8EF", fg="#087443")
        else:
            self.ai_status.config(bg="#FDECEC", fg="#B42318")
        if hasattr(self, "mood_status"):
            self.mood_status.config(
                text=f"Mood: {self.current_emotion or 'neutral'}"
                if not self.current_crisis_flag else "Support check"
            )

    def _create_video_overview(self, storyboard):
        """Render the generated storyboard to MP4 using the bundled imageio-ffmpeg binary."""
        try:
            from PIL import Image, ImageDraw, ImageFont
            asset_dir = topic_assets_dir(self.current_topic)
            video_dir = os.path.join(asset_dir, "video_overview")
            os.makedirs(video_dir, exist_ok=True)
            scenes = re.split(r"\n\s*(?=SCENE\s*\d+)|\n\s*(?=Scene\s*\d+)", storyboard, flags=re.I)
            scenes = [item.strip() for item in scenes if item.strip()] or [storyboard]
            try:
                font = ImageFont.truetype("arial.ttf", 30)
                small = ImageFont.truetype("arial.ttf", 22)
            except Exception:
                font = ImageFont.load_default()
                small = ImageFont.load_default()
            image_paths = []
            for index, scene in enumerate(scenes[:12], 1):
                image = Image.new("RGB", (1280, 720), "white")
                draw = ImageDraw.Draw(image)
                draw.text((55, 45), f"ArkaAI Video Overview — Scene {index}", fill="black", font=font)
                y = 120
                for line in textwrap.wrap(scene, width=70)[:22]:
                    draw.text((60, y), line, fill="black", font=small)
                    y += 27
                path = os.path.join(video_dir, f"scene_{index:02d}.png")
                image.save(path)
                image_paths.append(path)

            ffmpeg = shutil.which("ffmpeg") or shutil.which("ffmpeg.exe")
            if not ffmpeg:
                try:
                    import imageio_ffmpeg
                    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
                except Exception:
                    ffmpeg = None
            if not ffmpeg:
                return None

            output = os.path.join(
                asset_dir,
                f"Video_Overview_{datetime.now().strftime('%Y%m%d_%H%M%S')}.mp4",
            )
            subprocess.run(
                [ffmpeg, "-y", "-framerate", "1/4", "-i",
                 os.path.join(video_dir, "scene_%02d.png"),
                 "-c:v", "libx264", "-pix_fmt", "yuv420p", output],
                stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                timeout=240, check=False,
            )
            return output if os.path.exists(output) else None
        except Exception as exc:
            print("Video overview rendering skipped:", exc)
            return None

    def route_message(self, text):
        lower = text.lower().strip()
        if lower in {"profile", "my profile", "learning profile", "my learning profile", "memory"}:
            self.show_learning_profile()
            return
        if lower in {"search memory", "remember this", "what do you remember about me"}:
            memory = _semantic_memory_context(text, topic=self.current_topic, k=8)
            self.add_message(memory or "I don’t have enough saved semantic memory yet.", "ai")
            return
        super().route_message(text)


if __name__ == "__main__":
    root = tk.Tk()
    app = ArkaAIApp(root)
    root.mainloop()
